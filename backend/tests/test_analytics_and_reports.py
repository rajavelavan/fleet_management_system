import io
import uuid

import openpyxl
import pytest
from httpx import AsyncClient

from app.database import AsyncSessionLocal
from app.models.vehicle import Vehicle, VehicleStatus, VehicleType
from app.models.vehicle_location import VehicleLocation
from app.tasks.scheduled_tasks import (
    _async_daily_fleet_summary,
    _async_mark_unresponsive_vehicles,
)


@pytest.mark.asyncio
async def test_kpi_summary_endpoint(client: AsyncClient):
    login_res = await client.post(
        "/api/v1/auth/login",
        json={"email": "admin@fleet.local", "password": "AdminPassword123!"},
    )
    assert login_res.status_code == 200
    token = login_res.json()["access_token"]
    headers = {"Authorization": f"Bearer {token}"}

    response = await client.get("/api/v1/analytics/kpi-summary", headers=headers)
    assert response.status_code == 200
    data = response.json()
    assert "fleet_status" in data
    assert "delivery_status" in data
    assert "active_alerts_count" in data
    assert "on_time_delivery_rate" in data
    assert "total_trips_completed" in data
    assert isinstance(data["fleet_status"]["total"], int)
    assert isinstance(data["delivery_status"]["total"], int)


@pytest.mark.asyncio
async def test_deliveries_trend_endpoint(client: AsyncClient):
    login_res = await client.post(
        "/api/v1/auth/login",
        json={"email": "admin@fleet.local", "password": "AdminPassword123!"},
    )
    assert login_res.status_code == 200
    token = login_res.json()["access_token"]
    headers = {"Authorization": f"Bearer {token}"}

    response = await client.get("/api/v1/analytics/deliveries-trend?days=7", headers=headers)
    assert response.status_code == 200
    data = response.json()
    assert data["days"] == 7
    assert len(data["data"]) == 7
    assert "date" in data["data"][0]
    assert "delivered" in data["data"][0]
    assert "failed" in data["data"][0]
    assert "total" in data["data"][0]


@pytest.mark.asyncio
async def test_alerts_breakdown_endpoint(client: AsyncClient):
    login_res = await client.post(
        "/api/v1/auth/login",
        json={"email": "admin@fleet.local", "password": "AdminPassword123!"},
    )
    assert login_res.status_code == 200
    token = login_res.json()["access_token"]
    headers = {"Authorization": f"Bearer {token}"}

    response = await client.get("/api/v1/analytics/alerts-breakdown", headers=headers)
    assert response.status_code == 200
    data = response.json()
    assert "total_alerts" in data
    assert "breakdown" in data
    assert isinstance(data["breakdown"], list)


@pytest.mark.asyncio
async def test_vehicle_trip_summary_endpoint(client: AsyncClient):
    login_res = await client.post(
        "/api/v1/auth/login",
        json={"email": "admin@fleet.local", "password": "AdminPassword123!"},
    )
    assert login_res.status_code == 200
    token = login_res.json()["access_token"]
    headers = {"Authorization": f"Bearer {token}"}

    plate = f"TRIP-{uuid.uuid4().hex[:6].upper()}"
    async with AsyncSessionLocal() as session:
        v = Vehicle(
            plate_number=plate,
            make="Volvo",
            model="FH16",
            year=2024,
            type=VehicleType.TRUCK,
            status=VehicleStatus.ACTIVE,
            fuel_capacity=300.0,
            current_fuel=200.0,
        )
        session.add(v)
        await session.commit()
        await session.refresh(v)
        v_id = v.id

        # Ingest 2 location telemetry points
        loc1 = VehicleLocation(
            vehicle_id=v_id,
            location="SRID=4326;POINT(13.4050 52.5200)",
            speed=40.0,
        )
        loc2 = VehicleLocation(
            vehicle_id=v_id,
            location="SRID=4326;POINT(13.4150 52.5250)",
            speed=55.0,
        )
        session.add_all([loc1, loc2])
        await session.commit()

    # Query trip summary via API
    response = await client.get(f"/api/v1/vehicles/{v_id}/trip-summary", headers=headers)
    assert response.status_code == 200
    data = response.json()
    assert data["vehicle_id"] == str(v_id)
    assert data["plate_number"] == plate
    assert data["total_distance_km"] > 0.0
    assert data["max_speed_kmh"] == 55.0
    assert data["avg_speed_kmh"] == 47.5


@pytest.mark.asyncio
async def test_export_deliveries_xlsx(client: AsyncClient):
    login_res = await client.post(
        "/api/v1/auth/login",
        json={"email": "admin@fleet.local", "password": "AdminPassword123!"},
    )
    assert login_res.status_code == 200
    token = login_res.json()["access_token"]
    headers = {"Authorization": f"Bearer {token}"}

    response = await client.get("/api/v1/reports/deliveries/export?format=xlsx", headers=headers)
    assert response.status_code == 200
    assert response.headers["content-type"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    assert "attachment; filename=" in response.headers["content-disposition"]

    # Load xlsx bytes with openpyxl to verify valid format
    wb = openpyxl.load_workbook(io.BytesIO(response.content))
    ws = wb.active
    assert ws.title == "Deliveries Report"
    assert "DELIVERIES REPORT" in ws["A1"].value
    assert ws["A4"].value == "Order #"


@pytest.mark.asyncio
async def test_export_trips_xlsx(client: AsyncClient):
    login_res = await client.post(
        "/api/v1/auth/login",
        json={"email": "admin@fleet.local", "password": "AdminPassword123!"},
    )
    assert login_res.status_code == 200
    token = login_res.json()["access_token"]
    headers = {"Authorization": f"Bearer {token}"}

    response = await client.get("/api/v1/reports/trips/export?format=xlsx", headers=headers)
    assert response.status_code == 200
    assert response.headers["content-type"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    assert "attachment; filename=" in response.headers["content-disposition"]

    # Verify openpyxl can read
    wb = openpyxl.load_workbook(io.BytesIO(response.content))
    ws = wb.active
    assert ws.title == "Vehicle Trips Summary"
    assert "VEHICLE TRIPS" in ws["A1"].value
    assert ws["A4"].value == "Plate Number"


@pytest.mark.asyncio
async def test_scheduled_mark_unresponsive_vehicles_logic(db_session):
    plate = f"OFFLINE-{uuid.uuid4().hex[:6].upper()}"
    v = Vehicle(
        plate_number=plate,
        make="Ford",
        model="Transit",
        year=2023,
        type=VehicleType.VAN,
        status=VehicleStatus.ACTIVE,
        last_seen_at=None,
    )
    db_session.add(v)
    await db_session.commit()
    await db_session.refresh(v)

    # Run background check function with custom session or directly
    res = await _async_mark_unresponsive_vehicles(threshold_minutes=15, session=db_session)
    assert res["marked_offline_count"] >= 1

    await db_session.refresh(v)
    assert v.status == VehicleStatus.OFFLINE


@pytest.mark.asyncio
async def test_scheduled_daily_fleet_summary_digest(db_session):
    res = await _async_daily_fleet_summary(session=db_session)
    assert "delivered_count" in res
    assert "active_vehicles" in res
    assert "alerts_count" in res
